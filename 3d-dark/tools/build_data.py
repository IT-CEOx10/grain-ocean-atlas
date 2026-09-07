# -*- coding: utf-8 -*-
"""Сборка data/export.json из исходного xlsx.

Что делает:
  - читает лист «по странам» (иерархия задана отступом ячейки A:
    0 — группа продукции, 1 — продукт, 2 — страна);
  - берёт годы из строки заголовков (2014..2025);
  - подтягивает координаты и iso-коды из data/countries_ru.json;
  - пишет data/export.json;
  - сверяет суммы с листом «по годам» и печатает таблицу сравнения.

Запуск:  python3 tools/build_data.py
Если xlsx заменили — просто запустить скрипт заново, код править не нужно.
"""
from __future__ import print_function

import datetime
import glob
import json
import os
import sys

try:
    import openpyxl
except ImportError:
    print("Нужен openpyxl:  pip3 install openpyxl")
    sys.exit(1)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(ROOT, "data")
SHEET_COUNTRIES = "по странам"
SHEET_YEARS = "по годам"
TOTAL_ROW = "общий итог"          # строка-итог, её пропускаем
YEARS_LIMIT = (2014, 2025)        # таймлайн приложения
ROUND = 3


def nrm(value):
    """Нормализация текста: убираем неразрывные пробелы и лишние пробелы."""
    if value is None:
        return ""
    return " ".join(str(value).replace(u"\xa0", " ").split())


def find_xlsx():
    files = sorted(glob.glob(os.path.join(DATA_DIR, "*.xlsx")))
    files = [f for f in files if not os.path.basename(f).startswith("~$")]
    if not files:
        die("В папке data/ нет ни одного .xlsx с исходными данными.")
    if len(files) > 1:
        print("В data/ несколько xlsx, беру самый свежий по дате изменения:")
        for f in files:
            print("   " + os.path.basename(f))
        files.sort(key=os.path.getmtime)
    return files[-1]


def die(msg, details=None):
    print("")
    print("ОШИБКА: " + msg)
    if details:
        for d in details:
            print("   " + d)
    print("")
    sys.exit(1)


def read_header_years(ws):
    """Годы из строки заголовков. Ищем строку, где в B стоит год."""
    for row in range(1, 12):
        vals = []
        for col in range(2, ws.max_column + 1):
            raw = nrm(ws.cell(row, col).value)
            if raw[:4].isdigit() and 1990 < int(raw[:4]) < 2100:
                vals.append((col, int(raw[:4])))
            else:
                vals.append(None)
        if vals and vals[0]:
            return row, [v for v in vals if v]
    die("Не нашёл строку с годами в заголовке листа «%s»." % ws.title)


def num(cell_value):
    if cell_value is None:
        return 0.0
    if isinstance(cell_value, (int, float)):
        return float(cell_value)
    txt = nrm(cell_value).replace(",", ".").replace(" ", "")
    try:
        return float(txt)
    except ValueError:
        return 0.0


def parse_countries_sheet(ws, header_row, year_cols):
    """Возвращает: список групп, список продуктов, данные по странам.

    products: [{"n": имя, "g": индекс группы}]
    rows: {страна: {индекс продукта: {год: объём}}}
    group_totals: {год: {группа: объём}} — из строк-групп (indent 0)
    """
    groups = []
    group_index = {}
    products = []
    product_index = {}
    rows = {}
    group_totals = {}
    cur_group = None
    cur_product = None
    skipping_total = False

    for r in range(header_row + 1, ws.max_row + 1):
        cell = ws.cell(r, 1)
        name = nrm(cell.value)
        if not name:
            continue
        indent = int(cell.alignment.indent or 0)

        if indent == 0:
            if name.lower() == TOTAL_ROW:
                skipping_total = True
                cur_group = cur_product = None
                continue
            skipping_total = False
            if name not in group_index:
                group_index[name] = len(groups)
                groups.append(name)
            cur_group = group_index[name]
            cur_product = None
            for col, year in year_cols:
                group_totals.setdefault(year, {})[name] = \
                    group_totals.setdefault(year, {}).get(name, 0.0) + num(ws.cell(r, col).value)
            continue

        if skipping_total or cur_group is None:
            continue

        if indent == 1:
            key = (cur_group, name)
            if key not in product_index:
                product_index[key] = len(products)
                products.append({"n": name, "g": cur_group})
            cur_product = product_index[key]
            continue

        # indent >= 2 — страна под текущим продуктом
        if cur_product is None:
            continue
        bucket = rows.setdefault(name, {}).setdefault(cur_product, {})
        for col, year in year_cols:
            v = num(ws.cell(r, col).value)
            if v:
                bucket[year] = bucket.get(year, 0.0) + v

    return groups, products, rows, group_totals


def parse_years_sheet(ws):
    """Итоги по годам с контрольного листа: {год: значение строки «Общий итог»}."""
    header_row, year_cols = read_header_years(ws)
    totals = {}
    for r in range(header_row + 1, ws.max_row + 1):
        if nrm(ws.cell(r, 1).value).lower() == TOTAL_ROW:
            for col, year in year_cols:
                totals[year] = num(ws.cell(r, col).value)
            break
    return totals


def main():
    xlsx = find_xlsx()
    print("Источник: %s" % os.path.basename(xlsx))

    ref_path = os.path.join(DATA_DIR, "countries_ru.json")
    if not os.path.exists(ref_path):
        die("Нет справочника data/countries_ru.json.")
    ref = json.load(open(ref_path, encoding="utf-8"))
    ref = {nrm(k): v for k, v in ref.items()}

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    if SHEET_COUNTRIES not in wb.sheetnames:
        die("В файле нет листа «%s». Есть: %s" % (SHEET_COUNTRIES, ", ".join(wb.sheetnames)))
    ws = wb[SHEET_COUNTRIES]

    header_row, year_cols = read_header_years(ws)
    year_cols = [(c, y) for c, y in year_cols if YEARS_LIMIT[0] <= y <= YEARS_LIMIT[1]]
    years = [y for _, y in year_cols]
    print("Годы: %s" % ", ".join(str(y) for y in years))

    groups, products, rows, group_totals = parse_countries_sheet(ws, header_row, year_cols)
    print("Групп продукции: %d, продуктов: %d, стран в файле: %d"
          % (len(groups), len(products), len(rows)))

    missing = sorted(n for n in rows if n not in ref)
    if missing:
        die("В справочнике data/countries_ru.json нет %d стран. Добавьте их "
            "(iso из data/geo/countries-110m.json или null, обязательно lat/lon и en):" % len(missing),
            missing)

    # --- собираем страны -----------------------------------------------------
    countries = []
    for name in sorted(rows):
        info = ref[name]
        per_year = {}
        totals = []
        for y in years:
            items = []
            exact = 0.0
            for pi, by_year in rows[name].items():
                v = by_year.get(y, 0.0)
                if v:
                    items.append([pi, round(v, ROUND)])
                    exact += v          # сумма считается до округления позиций
            items.sort(key=lambda it: -it[1])
            if items:
                per_year[str(y)] = items
            totals.append(round(exact, ROUND))
        countries.append({
            "name": name,
            "en": info.get("en") or name,
            "iso": info.get("iso"),
            "lat": info["lat"],
            "lon": info["lon"],
            "totals": totals,
            "years": per_year,
        })

    # --- сводка по годам -----------------------------------------------------
    summary = {}
    for i, y in enumerate(years):
        vals = [(c["name"], c["totals"][i]) for c in countries if c["totals"][i] > 0]
        vals.sort(key=lambda t: -t[1])
        used_groups = set()
        for c in countries:
            for pi, v in c["years"].get(str(y), []):
                used_groups.add(products[pi]["g"])
        summary[str(y)] = {
            "total": round(sum(v for _, v in vals), ROUND),
            "countries": len(vals),
            "groups": len(used_groups),
            "top": [{"name": n, "value": v} for n, v in vals[:3]],
        }

    out = {
        "generated": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "source": os.path.basename(xlsx),
        "unit": "тыс. тонн",
        "years": years,
        "groups": groups,
        "products": products,
        "countries": countries,
        "summary": summary,
    }

    out_path = os.path.join(DATA_DIR, "export.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
    size = os.path.getsize(out_path) / 1024.0
    print("Записано: %s (%.0f КБ)" % (os.path.relpath(out_path, ROOT), size))

    # --- сверка --------------------------------------------------------------
    control = parse_years_sheet(wb[SHEET_YEARS]) if SHEET_YEARS in wb.sheetnames else {}
    print("")
    print("Сверка сумм, тыс. тонн")
    print("%-6s %14s %14s %14s %10s" % ("год", "по странам", "группы", "лист по годам", "расхожд."))
    bad = 0
    for i, y in enumerate(years):
        by_countries = summary[str(y)]["total"]
        by_groups = round(sum(group_totals.get(y, {}).values()), ROUND)
        ctrl = control.get(y)
        diff = (by_countries - ctrl) if ctrl is not None else 0.0
        # допуск: суммируем тысячи значений, округлённых до 3 знаков
        if ctrl is not None and abs(diff) > 0.5:
            bad += 1
        print("%-6d %14.3f %14.3f %14s %10.3f"
              % (y, by_countries, by_groups,
                 ("%.3f" % ctrl) if ctrl is not None else "—", diff))
    print("")
    no_iso = [c["name"] for c in countries if not c["iso"]]
    print("Стран без iso в world-atlas (рисуем только маркер): %d — %s"
          % (len(no_iso), ", ".join(no_iso)))
    if bad:
        print("ВНИМАНИЕ: расхождение больше 0.5 тыс. т в %d годах — проверьте исходный файл." % bad)
    else:
        print("Все годы сходятся с листом «по годам» (в пределах округления).")


if __name__ == "__main__":
    main()
